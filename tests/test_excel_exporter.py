from io import BytesIO

import openpyxl

from app.export.excel_exporter import export_excel
from tests.factories import build_sample_deliverable


def test_export_excel_produces_expected_sheets():
    deliverable = build_sample_deliverable()
    workbook_bytes = export_excel(deliverable)

    wb = openpyxl.load_workbook(BytesIO(workbook_bytes))
    assert set(wb.sheetnames) == {
        "Test Cases",
        "Test Scenarios",
        "Impact Analysis",
        "Risks",
        "Automation",
        "Traceability",
    }

    test_cases_sheet = wb["Test Cases"]
    header = [cell.value for cell in next(test_cases_sheet.iter_rows(min_row=1, max_row=1))]
    assert "Test Case ID" in header
    assert "Expected Results" in header

    ids_column = header.index("Test Case ID")
    ids = [row[ids_column].value for row in test_cases_sheet.iter_rows(min_row=2)]
    assert "TC-001" in ids
    assert "TC-002" in ids
