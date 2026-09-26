"""Extracts plain text from the supported requirement document formats.

Supported: .txt, .md, .csv, .docx, .pdf, .xlsx/.xls, and common image formats
(best-effort OCR if pytesseract + Pillow + a system tesseract binary are
available; otherwise a clear placeholder is returned).
"""
from __future__ import annotations

from pathlib import Path


class UnsupportedFileTypeError(ValueError):
    pass


def _read_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def _read_docx(path: Path) -> str:
    import docx  # python-docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages)


def _read_excel(path: Path) -> str:
    import pandas as pd

    sheets = pd.read_excel(str(path), sheet_name=None, dtype=str)
    chunks = []
    for name, df in sheets.items():
        df = df.fillna("")
        chunks.append(f"--- Sheet: {name} ---")
        chunks.append(df.to_csv(index=False))
    return "\n".join(chunks)


def _read_csv(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def _read_image(path: Path) -> str:
    try:
        import pytesseract
        from PIL import Image

        text = pytesseract.image_to_string(Image.open(path))
        text = text.strip()
        if text:
            return text
    except Exception:
        pass
    return (
        f"[Image file '{path.name}' could not be OCR'd automatically. "
        "Please describe its content (e.g. wireframe/screenshot details) as additional text input.]"
    )


_HANDLERS = {
    ".txt": _read_txt,
    ".md": _read_txt,
    ".csv": _read_csv,
    ".docx": _read_docx,
    ".pdf": _read_pdf,
    ".xlsx": _read_excel,
    ".xls": _read_excel,
    ".png": _read_image,
    ".jpg": _read_image,
    ".jpeg": _read_image,
    ".gif": _read_image,
    ".bmp": _read_image,
}


def extract_text(path: str | Path) -> str:
    path = Path(path)
    suffix = path.suffix.lower()
    handler = _HANDLERS.get(suffix)
    if handler is None:
        raise UnsupportedFileTypeError(
            f"Unsupported file type '{suffix}'. Supported: {', '.join(sorted(_HANDLERS))}"
        )
    return handler(path)
